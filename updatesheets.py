from __future__ import annotations

import argparse
from datetime import datetime
from pathlib import Path

import requests

DEFAULT_WORKBOOK = Path("Identified_Cards") / "automatic_runs.xlsx"


def get_cad_price(prices: dict, usd_to_cad: float = 1.37, eur_to_cad: float = 1.57, default_price=None):
    """Convert USD/EUR price fields to CAD and return (value, source)."""
    if prices.get("usd"):
        try:
            return round(float(prices["usd"]) * usd_to_cad, 2), "usd"
        except Exception:
            pass
    if prices.get("eur"):
        try:
            return round(float(prices["eur"]) * eur_to_cad, 2), "eur"
        except Exception:
            pass
    if default_price is not None:
        return default_price, "default"
    return None, None


def fetch_prices_from_scryfall(api_url: str) -> dict:
    if not api_url:
        return {}
    try:
        response = requests.get(api_url, timeout=15)
    except Exception:
        return {}
    if response.status_code != 200:
        return {}
    try:
        payload = response.json()
    except Exception:
        return {}
    return payload.get("prices", {}) or {}


def update_workbook_prices(workbook_path: Path) -> tuple[int, int]:
    try:
        import openpyxl  # type: ignore
    except Exception as exc:
        raise RuntimeError("openpyxl is required to update Excel sheets.") from exc

    load_workbook = openpyxl.load_workbook
    wb = load_workbook(str(workbook_path))
    updated_rows = 0
    visited_rows = 0

    for ws in wb.worksheets:
        if ws.max_row < 2:
            continue

        # Expected columns:
        # A card name, B camera image link, C scryfall api link, D current price cad, E confidence
        for row_idx in range(2, ws.max_row + 1):
            visited_rows += 1
            api_cell = ws.cell(row=row_idx, column=3)
            price_cell = ws.cell(row=row_idx, column=4)

            api_url = str(api_cell.value or "").strip()
            if not api_url:
                continue

            prices = fetch_prices_from_scryfall(api_url)
            cad_price, _source = get_cad_price(prices, default_price=price_cell.value)
            if cad_price is None:
                continue

            old_val = price_cell.value
            price_cell.value = cad_price
            if old_val != cad_price:
                updated_rows += 1

    wb.save(str(workbook_path))
    wb.close()
    return visited_rows, updated_rows


def main():
    parser = argparse.ArgumentParser(description="Update card prices in automatic run workbook sheets.")
    parser.add_argument(
        "--workbook",
        type=str,
        default=str(DEFAULT_WORKBOOK),
        help="Path to the workbook to update.",
    )
    args = parser.parse_args()

    workbook_path = Path(args.workbook)
    if not workbook_path.exists():
        print(f"Workbook not found: {workbook_path}")
        return

    started = datetime.now()
    visited, updated = update_workbook_prices(workbook_path)
    finished = datetime.now()

    print(f"Workbook updated: {workbook_path}")
    print(f"Rows visited: {visited}")
    print(f"Rows changed: {updated}")
    print(f"Started: {started}")
    print(f"Finished: {finished}")


if __name__ == "__main__":
    main()
